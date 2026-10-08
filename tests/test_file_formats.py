"""
Tests for PDF and DOCX parsers with actual binary documents.
"""

import pytest
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document

from src.parsers.pdf_parser import PDFParser
from src.parsers.docx_parser import DOCXParser


def test_pdf_parser_real_pdf(tmp_path):
    """Test generating a real PDF and reading text with PDFParser."""
    pdf_path = tmp_path / "sample.pdf"
    doc = SimpleDocTemplate(str(pdf_path), pagesize=letter)
    styles = getSampleStyleSheet()
    story = [Paragraph("Candidate: Vikram Seth. Python FastAPI Developer.", styles["Normal"])]
    doc.build(story)

    parser = PDFParser()
    text = parser.extract_text(pdf_path)
    assert "Vikram Seth" in text
    assert "FastAPI" in text


def test_docx_parser_real_docx(tmp_path):
    """Test generating a real DOCX and reading text with DOCXParser."""
    docx_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Arjun Patel", level=1)
    doc.add_paragraph("Skills: Python, LangGraph, PostgreSQL, Docker")
    doc.save(str(docx_path))

    parser = DOCXParser()
    text = parser.extract_text(docx_path)
    assert "Arjun Patel" in text
    assert "LangGraph" in text


def test_corrupt_pdf_fails_gracefully(tmp_path):
    """Test that corrupt PDF raises ValueError without unhandled crash."""
    corrupt_path = tmp_path / "corrupt.pdf"
    corrupt_path.write_bytes(b"NOT A VALID PDF")

    parser = PDFParser()
    with pytest.raises(ValueError) as excinfo:
        parser.extract_text(corrupt_path)
    assert "corrupt" in str(excinfo.value).lower() or "failed" in str(excinfo.value).lower()
