"""
DOCX resume parser implementation.
"""

from pathlib import Path
import docx
from src.parsers.base import BaseParser


class DOCXParser(BaseParser):
    """Extracts text from DOCX documents."""

    def extract_text(self, file_path: Path) -> str:
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        if file_path.stat().st_size == 0:
            raise ValueError(f"File is empty (0 bytes): {file_path.name}")

        try:
            doc = docx.Document(str(file_path))
            chunks = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    chunks.append(text)

            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        chunks.append(row_text)

            full_text = "\n".join(chunks).strip()
            if not full_text:
                raise ValueError(f"No extractable text found in DOCX: {file_path.name}")
            return full_text
        except Exception as e:
            if isinstance(e, ValueError):
                raise
            raise ValueError(f"Failed to parse DOCX {file_path.name}: {str(e)}")
