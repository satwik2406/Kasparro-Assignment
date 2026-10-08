"""
Synthetic Resume Generator.
Generates 50 realistic, diverse candidate resumes in PDF, DOCX, and TXT formats
to benchmark the AI Resume Screening & Ranking System.
"""

import os
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from docx import Document


RESUMES_DIR = Path("./resumes")
RESUMES_DIR.mkdir(parents=True, exist_ok=True)


CANDIDATES = [
    # Top Tier Agentic & AI Engineers (Pass hard filter, High scores 75-95)
    {
        "name": "Asha Rao",
        "email": "asha.rao@example.com",
        "phone": "+1-555-0101",
        "github": "https://github.com/asharao",
        "format": "pdf",
        "skills": "Python, FastAPI, AsyncIO, LangGraph, Qdrant, RAG, Tool Calling, Multi-Agent, Ragas, PostgreSQL, Redis, Docker, GCP, pytest",
        "experience": "AI Engineering Intern at Agentic Labs. Built multi-agent research assistant using LangGraph with state persistence and tool-calling capabilities. Deployed on GCP using Docker and Cloud Run.",
        "projects": "Autonomous Workflow Orchestrator: Built production RAG with hybrid search (Qdrant + BM25) and automated evaluation using Ragas. Integrated async FastAPI backend with Redis caching and PostgreSQL storage.",
        "education": "B.S. in Computer Science, GPA: 3.9/4.0"
    },
    {
        "name": "Vikram Seth",
        "email": "vikram.seth@example.com",
        "phone": "+1-555-0102",
        "github": "https://github.com/vikramseth",
        "format": "pdf",
        "skills": "Python, LangGraph, CrewAI, FastAPI, ChromaDB, Vector Search, PostgreSQL, Docker, GCP, PyTorch, pytest",
        "experience": "Software Intern at Cognitive AI. Engineered a multi-agent triage system with LangGraph and ChromaDB. Optimized retrieval latency by 40%.",
        "projects": "Agentic Code Reviewer: Built an autonomous agent with function calling to lint, test, and comment on PRs. Full FastAPI REST API containerized with Docker on Google Cloud.",
        "education": "M.S. in Computer Science"
    },
    {
        "name": "Maya Lin",
        "email": "maya.lin@example.com",
        "phone": "+1-555-0103",
        "github": "https://github.com/mayalin",
        "format": "pdf",
        "skills": "Python, FastAPI, LangChain, Pinecone, RAG, Tool Calling, Redis, Docker, GCP, React, TypeScript",
        "experience": "SWE Intern at DataFlow. Built end-to-end RAG pipeline ingesting 100k financial documents into Pinecone vector database with reranking.",
        "projects": "Enterprise Doc QA: Full stack system with FastAPI backend, Redis caching, Docker deployment, and Next.js/React frontend. Evaluated RAG precision and recall.",
        "education": "B.S. in Computer Engineering"
    },
    {
        "name": "Arjun Patel",
        "email": "arjun.patel@example.com",
        "phone": "+1-555-0104",
        "github": "https://github.com/arjunpatel",
        "format": "docx",
        "skills": "Python, AsyncIO, FastAPI, LlamaIndex, Qdrant, Multi-Agent, SQLAlchemy, PostgreSQL, Docker, AWS",
        "experience": "Backend Developer Intern. Architected asynchronous event-driven pipelines using FastAPI, PostgreSQL, and RabbitMQ.",
        "projects": "Multi-Agent Research Engine: Coordinated specialist agents using LlamaIndex and Qdrant. Includes circuit-breaker error handling and pytest test coverage.",
        "education": "B.Tech in Computer Science"
    },
    {
        "name": "Sophia Chen",
        "email": "sophia.chen@example.com",
        "phone": "+1-555-0105",
        "github": "https://github.com/sophiachen",
        "format": "pdf",
        "skills": "Python, PyTorch, LangGraph, RAG, Embeddings, Docker, GCP, PostgreSQL, Redis, CI/CD",
        "experience": "Research Assistant. Trained embeddings and integrated vector databases for medical literature retrieval.",
        "projects": "Clinical Assistant: Stateful multi-agent diagnostic helper using LangGraph, Docker containerization, and GCP deployment with automated tests.",
        "education": "B.S. in Data Science & CS"
    },
    {
        "name": "Rohan Gupta",
        "email": "rohan.gupta@example.com",
        "phone": "+1-555-0106",
        "github": "https://github.com/rohangupta",
        "format": "pdf",
        "skills": "Python, FastAPI, LangChain, FAISS, Docker, PostgreSQL, Redis, Celery, Linux",
        "experience": "Backend Engineer Intern. Developed high-throughput microservices using FastAPI, Redis caching, and Celery workers.",
        "projects": "Semantic Knowledge Engine: RAG search across corporate wiki using FAISS vector search, FastAPI, and Dockerized deployment.",
        "education": "B.E. in Information Technology"
    },
    {
        "name": "Elena Rossi",
        "email": "elena.rossi@example.com",
        "phone": "+1-555-0107",
        "github": "https://github.com/elenarossi",
        "format": "docx",
        "skills": "Python, FastAPI, LangGraph, Function Calling, PostgreSQL, Redis, Docker, GCP, pytest",
        "experience": "Software Engineer Intern. Implemented tool-calling agent pipelines for automated SQL querying and schema validation.",
        "projects": "Autonomous Data Analyst: LangGraph stateful agent orchestrating schema inspection, SQL generation, and evaluation with pytest.",
        "education": "B.S. in Software Engineering"
    },
    {
        "name": "Carlos Mendez",
        "email": "carlos.mendez@example.com",
        "phone": "+1-555-0108",
        "github": "https://github.com/carlosmendez",
        "format": "pdf",
        "skills": "Python, FastAPI, LlamaIndex, ChromaDB, RAG, Docker, GCP, PostgreSQL",
        "experience": "AI Intern at HealthTech. Built HIPAA-compliant RAG pipeline using ChromaDB and FastAPI.",
        "projects": "Intelligent Triage System: Vector search over medical records with hybrid retrieval and GCP deployment.",
        "education": "B.S. in Computer Science"
    },
    {
        "name": "Priya Sharma",
        "email": "priya.sharma@example.com",
        "phone": "+1-555-0109",
        "github": "https://github.com/priyasharma",
        "format": "pdf",
        "skills": "Python, LangGraph, CrewAI, FastAPI, PostgreSQL, Redis, Docker, GCP, React",
        "experience": "SWE Intern. Engineered automated customer support agents using CrewAI and FastAPI.",
        "projects": "Collaborative Multi-Agent Support: Hierarchical agent team with fallback handling, Redis cache, and Docker deployment.",
        "education": "B.Tech in Computer Science"
    },
    {
        "name": "David Kim",
        "email": "david.kim@example.com",
        "phone": "+1-555-0110",
        "github": "https://github.com/davidkim",
        "format": "pdf",
        "skills": "Python, FastAPI, RAG, Weaviate, Tool Calling, Docker, AWS, PostgreSQL, pytest",
        "experience": "AI Research Intern. Built RAG benchmark comparing Weaviate and Pinecone with automated evaluation.",
        "projects": "Legal Contract Auditor: RAG system extracting clauses with tool calling and automated test suites in pytest.",
        "education": "B.S. in Computer Science"
    },

    # Solid RAG & LLM Engineers (Pass hard filter, Scores 60-75)
    {
        "name": "Liam O'Connor",
        "email": "liam.oconnor@example.com",
        "phone": "+1-555-0111",
        "github": "https://github.com/liamoconnor",
        "format": "pdf",
        "skills": "Python, FastAPI, LlamaIndex, Pinecone, Vector Search, PostgreSQL, Docker",
        "experience": "Junior Software Developer. Built REST APIs in FastAPI with PostgreSQL.",
        "projects": "Document Search Engine: RAG application using LlamaIndex and Pinecone vector store with Docker deployment.",
        "education": "B.S. in Computer Science"
    },
    {
        "name": "Ananya Iyer",
        "email": "ananya.iyer@example.com",
        "phone": "+1-555-0112",
        "github": "https://github.com/ananyaiyer",
        "format": "docx",
        "skills": "Python, LangChain, ChromaDB, FastAPI, PostgreSQL, Git",
        "experience": "SWE Intern. Developed FastAPI microservices with PostgreSQL database integration.",
        "projects": "Academic Paper Q&A: RAG pipeline with chunking and Chroma vector search.",
        "education": "B.E. in Computer Science"
    },
    {
        "name": "Zachary Taylor",
        "email": "zach.taylor@example.com",
        "phone": "+1-555-0113",
        "github": "https://github.com/zachtaylor",
        "format": "pdf",
        "skills": "Python, Flask, LlamaIndex, FAISS, Docker, SQLite",
        "experience": "Software Engineering Fellow. Developed Python scripts and internal developer tooling.",
        "projects": "Internal Knowledge Base: Vector search pipeline using FAISS and LlamaIndex with Docker.",
        "education": "B.S. in Computer Engineering"
    },
    {
        "name": "Sneha Reddy",
        "email": "sneha.reddy@example.com",
        "phone": "+1-555-0114",
        "github": "https://github.com/snehareddy",
        "format": "pdf",
        "skills": "Python, FastAPI, RAG, ChromaDB, Embeddings, PostgreSQL, Docker",
        "experience": "Backend Intern. Implemented REST APIs and database models using SQLAlchemy and PostgreSQL.",
        "projects": "Customer FAQ RAG: Semantic search engine over product manuals using ChromaDB.",
        "education": "B.Tech in IT"
    },
    {
        "name": "Lucas Silva",
        "email": "lucas.silva@example.com",
        "phone": "+1-555-0115",
        "github": "https://github.com/lucassilva",
        "format": "txt",
        "skills": "Python, LangChain, Qdrant, Docker, PostgreSQL, REST API",
        "experience": "Developer Intern. Maintained Python microservices.",
        "projects": "RAG Document Extractor: Vector database retrieval system built with LangChain and Qdrant.",
        "education": "B.S. in Software Engineering"
    },
    {
        "name": "Chloe Martin",
        "email": "chloe.martin@example.com",
        "phone": "+1-555-0116",
        "github": "https://github.com/chloemartin",
        "format": "pdf",
        "skills": "Python, FastAPI, RAG, Pinecone, Docker, Linux",
        "experience": "Software Intern. Created data ingestion pipelines.",
        "projects": "News Summarizer & RAG: Vector search over news articles using Pinecone and FastAPI.",
        "education": "B.S. in Computer Science"
    },
    {
        "name": "Rahul Verma",
        "email": "rahul.verma@example.com",
        "phone": "+1-555-0117",
        "github": "https://github.com/rahulverma",
        "format": "pdf",
        "skills": "Python, Django, LangChain, ChromaDB, PostgreSQL, Docker",
        "experience": "Django Developer Intern. Built web dashboards.",
        "projects": "Company Policy RAG: Built QA chatbot over HR policies using LangChain and ChromaDB.",
        "education": "B.Tech in Computer Science"
    },
    {
        "name": "Fatima Zahra",
        "email": "fatima.zahra@example.com",
        "phone": "+1-555-0118",
        "github": "https://github.com/fatimazahra",
        "format": "pdf",
        "skills": "Python, FastAPI, LlamaIndex, PGVector, PostgreSQL, Docker",
        "experience": "AI Intern. Integrated PGVector into PostgreSQL for semantic retrieval.",
        "projects": "Semantic Catalog Search: Built hybrid keyword + vector retrieval service in FastAPI.",
        "education": "B.S. in Computer Science"
    },

    # Polyglots with Python + AI (Pass hard filter, Scores 55-75)
    {
        "name": "Hannah Polyglot",
        "email": "hannah.poly@example.com",
        "phone": "+1-555-0119",
        "github": "https://github.com/hannahpoly",
        "format": "pdf",
        "skills": "Java, Spring Boot, React, Python, FastAPI, LangChain, Qdrant, Docker",
        "experience": "Full Stack Intern. Built Spring Boot microservices and React web apps.",
        "projects": "AI Document Processing Service: Built Python FastAPI service using LangChain and Qdrant for semantic search alongside Java services.",
        "education": "B.S. in Computer Science"
    },
    {
        "name": "Oscar Fullstack",
        "email": "oscar.full@example.com",
        "phone": "+1-555-0120",
        "github": "https://github.com/oscarfull",
        "format": "docx",
        "skills": "TypeScript, React, Next.js, Python, FastAPI, LlamaIndex, Docker, PostgreSQL",
        "experience": "Full Stack Developer. Developed modern Next.js frontend applications.",
        "projects": "AI Research Workspace: Python FastAPI backend powering LlamaIndex vector retrieval with Next.js interface.",
        "education": "B.S. in Software Engineering"
    },
    {
        "name": "Zoe Hybrid",
        "email": "zoe.hybrid@example.com",
        "phone": "+1-555-0121",
        "github": "https://github.com/zoehybrid",
        "format": "txt",
        "skills": "Go, Python, LangGraph, FastAPI, Docker, Kubernetes, PostgreSQL",
        "experience": "Systems Intern. Built Golang network tools and Python agentic services.",
        "projects": "Multi-Agent System Orchestrator: Python LangGraph tool-calling agent with Docker deployment.",
        "education": "B.S. in Computer Engineering"
    },

    # Thin LLM Wrapper Candidates (Pass hard filter, PENALIZED 12 pts, Scores 35-50)
    {
        "name": "Kevin Zhang",
        "email": "kevin.zhang@example.com",
        "phone": "+1-555-0122",
        "github": "https://github.com/kevinzhang",
        "format": "pdf",
        "skills": "Python, Streamlit, OpenAI API, Flask, SQLite",
        "experience": "Undergraduate Student. Built simple web apps.",
        "projects": "AI Resume Feedback App: Simple Streamlit prompt wrapper calling OpenAI API prompt-response. No vector database, no retrieval, no tools.",
        "education": "B.S. in Informatics"
    },
    {
        "name": "Emily Watson",
        "email": "emily.watson@example.com",
        "phone": "+1-555-0123",
        "github": "https://github.com/emilywatson",
        "format": "pdf",
        "skills": "Python, Flask, OpenAI, ChatGPT API",
        "experience": "Student Developer.",
        "projects": "Recipe Generator: Basic prompt-response web page using Flask and OpenAI API call with no custom data processing.",
        "education": "B.A. in Computer Science"
    },
    {
        "name": "Tyler Brooks",
        "email": "tyler.brooks@example.com",
        "phone": "+1-555-0124",
        "github": "https://github.com/tylerbrooks",
        "format": "docx",
        "skills": "Python, Streamlit, OpenAI API",
        "experience": "Intern at local startup.",
        "projects": "Chat with Text: Streamlit app calling OpenAI API prompt-response directly without embeddings or vector search.",
        "education": "B.S. in Information Systems"
    },
    {
        "name": "Jessica Davis",
        "email": "jessica.davis@example.com",
        "phone": "+1-555-0125",
        "github": "https://github.com/jessicadavis",
        "format": "pdf",
        "skills": "Python, Flask, OpenAI API, SQLite",
        "experience": "Coding Bootcamp Graduate.",
        "projects": "AI Email Draft Assistant: Thin wrapper calling ChatGPT API via Flask backend with simple prompt string concatenation.",
        "education": "Coding Bootcamp Certificate"
    },
    {
        "name": "Noah Miller",
        "email": "noah.miller@example.com",
        "phone": "+1-555-0126",
        "github": "https://github.com/noahmiller",
        "format": "pdf",
        "skills": "Python, Streamlit, Gemini API, HTML",
        "experience": "Student Researcher.",
        "projects": "AI Story Generator: Streamlit UI directly forwarding user prompts to Gemini API with no backend architecture.",
        "education": "B.S. in Computer Science"
    },
    {
        "name": "Amanda White",
        "email": "amanda.white@example.com",
        "phone": "+1-555-0127",
        "github": "https://github.com/amandawhite",
        "format": "txt",
        "skills": "Python, OpenAI API, Tkinter",
        "experience": "Self-taught programmer.",
        "projects": "Desktop ChatGPT Client: Tkinter Python GUI wrapping OpenAI API completion endpoint.",
        "education": "High School Diploma"
    },

    # Tutorial / Cloned Projects (Pass hard filter, PENALIZED 8 pts, Scores 40-52)
    {
        "name": "Brandon Lee",
        "email": "brandon.lee@example.com",
        "phone": "+1-555-0128",
        "github": "https://github.com/brandonlee",
        "format": "pdf",
        "skills": "Python, LangChain, Streamlit, FAISS",
        "experience": "Student Developer.",
        "projects": "YouTube Video Summarizer: Followed YouTube tutorial project step by step to summarize video transcripts using LangChain.",
        "education": "B.S. in Computer Science"
    },
    {
        "name": "Samantha Green",
        "email": "sam.green@example.com",
        "phone": "+1-555-0129",
        "github": "https://github.com/samgreen",
        "format": "pdf",
        "skills": "Python, Streamlit, LangChain, Chroma",
        "experience": "CS Student.",
        "projects": "Chat with Multiple PDFs: Course project from Udemy tutorial cloned with boilerplate code.",
        "education": "B.S. in Computer Science"
    },
    {
        "name": "Daniel Clark",
        "email": "daniel.clark@example.com",
        "phone": "+1-555-0130",
        "github": "https://github.com/danielclark",
        "format": "docx",
        "skills": "Python, Flask, OpenAI, FAISS",
        "experience": "Junior Programmer.",
        "projects": "PDF Q&A Bot: Cloned from tutorial online without additional architecture or test suite.",
        "education": "B.S. in IT"
    },
    {
        "name": "Rachel Adams",
        "email": "rachel.adams@example.com",
        "phone": "+1-555-0131",
        "github": "https://github.com/racheladams",
        "format": "pdf",
        "skills": "Python, LangChain, Streamlit",
        "experience": "Developer.",
        "projects": "YouTube Transcriber: Followed online YouTube clone tutorial using LangChain.",
        "education": "B.S. in Computer Engineering"
    },

    # Pure Python Backend Engineers (REJECTED BY HARD FILTER: No AI/Agentic evidence)
    {
        "name": "Marcus Vance",
        "email": "marcus.vance@example.com",
        "phone": "+1-555-0132",
        "github": "https://github.com/marcusvance",
        "format": "pdf",
        "skills": "Python, Django, FastAPI, PostgreSQL, Redis, Celery, Docker, Linux, pytest",
        "experience": "Backend Engineer Intern. Built high-scale e-commerce REST APIs handling 5,000 req/sec with Redis caching and Celery queues.",
        "projects": "Distributed Order Management System: Event-driven architecture with PostgreSQL transactions and RabbitMQ.",
        "education": "B.S. in Computer Science"
    },
    {
        "name": "Nathan Drake",
        "email": "nathan.drake@example.com",
        "phone": "+1-555-0133",
        "github": "https://github.com/nathandrake",
        "format": "pdf",
        "skills": "Python, Flask, SQLAlchemy, MySQL, Docker, AWS",
        "experience": "Software Developer Intern. Built internal dashboard APIs and reporting pipelines.",
        "projects": "Fleet Management API: Microservice for vehicle tracking with geospatial database queries.",
        "education": "B.S. in Software Engineering"
    },
    {
        "name": "Olivia Stone",
        "email": "olivia.stone@example.com",
        "phone": "+1-555-0134",
        "github": "https://github.com/oliviastone",
        "format": "docx",
        "skills": "Python, FastAPI, AsyncIO, PostgreSQL, Redis, Kubernetes, Docker",
        "experience": "Infrastructure Intern. Maintained backend microservices and Kubernetes deployments.",
        "projects": "High Performance Payment Gateway: Low latency async FastAPI microservice with Redis lock mechanisms.",
        "education": "B.S. in Computer Science"
    },
    {
        "name": "Peter Parker",
        "email": "peter.parker@example.com",
        "phone": "+1-555-0135",
        "github": "https://github.com/peterparker",
        "format": "pdf",
        "skills": "Python, Django, PostgreSQL, Docker, Git",
        "experience": "Junior Backend Developer. Created CRUD endpoints and user authentication systems.",
        "projects": "School Management Portal: Django web application with PostgreSQL relational database.",
        "education": "B.S. in Computer Science"
    },
    {
        "name": "Rachel Green",
        "email": "rachel.green@example.com",
        "phone": "+1-555-0136",
        "github": "https://github.com/rachelgreen",
        "format": "pdf",
        "skills": "Python, Flask, SQLite, Pandas, NumPy",
        "experience": "Data Engineer Intern. Built ETL data pipelines using Pandas and SQLAlchemy.",
        "projects": "Sales Analytics Pipeline: Automated daily CSV extraction and cleaning into SQLite.",
        "education": "B.S. in Statistics"
    },
    {
        "name": "Thomas Anderson",
        "email": "thomas.anderson@example.com",
        "phone": "+1-555-0137",
        "github": "https://github.com/thomasanderson",
        "format": "txt",
        "skills": "Python, AsyncIO, FastAPI, PostgreSQL, Redis, Linux",
        "experience": "Systems Developer. Wrote asynchronous network crawlers and monitoring daemons.",
        "projects": "Async Server Heartbeat Monitor: AsyncIO daemon monitoring server metrics with alerts.",
        "education": "B.S. in Computer Science"
    },

    # Java / Spring Boot / C# Only (REJECTED BY HARD FILTER: No Python, No AI)
    {
        "name": "John Java",
        "email": "john.java@example.com",
        "phone": "+1-555-0138",
        "github": "https://github.com/johnjava",
        "format": "pdf",
        "skills": "Java, Spring Boot, Hibernate, MySQL, Maven, Docker",
        "experience": "Java Developer Intern. Built enterprise banking services in Spring Boot with Hibernate.",
        "projects": "Banking Core API: Spring Boot REST services with ACID transaction compliance.",
        "education": "B.S. in Computer Science"
    },
    {
        "name": "Sarah Spring",
        "email": "sarah.spring@example.com",
        "phone": "+1-555-0139",
        "github": "https://github.com/sarahspring",
        "format": "docx",
        "skills": "Java, Spring Boot, Microservices, Oracle DB, Kubernetes",
        "experience": "Junior Software Engineer. Maintained Spring Cloud microservices.",
        "projects": "Enterprise Inventory Management: Spring Boot backends with RabbitMQ.",
        "education": "B.E. in Computer Science"
    },
    {
        "name": "Mike DotNet",
        "email": "mike.dotnet@example.com",
        "phone": "+1-555-0140",
        "github": "https://github.com/mikedotnet",
        "format": "pdf",
        "skills": "C#, .NET Core, ASP.NET, SQL Server, Azure",
        "experience": "SWE Intern. Built web APIs in .NET Core hosted on Microsoft Azure.",
        "projects": "Hospital Record System: ASP.NET Core MVC web application with Entity Framework.",
        "education": "B.S. in Software Engineering"
    },
    {
        "name": "Chris Angular",
        "email": "chris.angular@example.com",
        "phone": "+1-555-0141",
        "github": "https://github.com/chrisangular",
        "format": "pdf",
        "skills": "Java, Spring Boot, Angular, TypeScript, PostgreSQL",
        "experience": "Full Stack Intern. Built Angular UI and Spring Boot backend.",
        "projects": "Corporate HR Portal: Angular frontend consuming Spring Boot REST APIs.",
        "education": "B.Tech in Information Technology"
    },
    {
        "name": "Dave Enterprise",
        "email": "dave.enterprise@example.com",
        "phone": "+1-555-0142",
        "github": "https://github.com/daveenterprise",
        "format": "pdf",
        "skills": "Java, Kotlin, Spring Boot, PostgreSQL, Docker",
        "experience": "Software Engineer Intern. Implemented microservices using Spring Boot and Kotlin.",
        "projects": "Billing Engine: Scalable billing service in Spring Boot.",
        "education": "B.S. in Computer Science"
    },

    # Frontend Only (REJECTED BY HARD FILTER: No Python, No AI)
    {
        "name": "Alice Frontend",
        "email": "alice.front@example.com",
        "phone": "+1-555-0143",
        "github": "https://github.com/alicefront",
        "format": "pdf",
        "skills": "JavaScript, React, Next.js, HTML, CSS, Tailwind CSS, TypeScript",
        "experience": "Frontend Intern. Built responsive UI components in React and Next.js.",
        "projects": "E-Commerce Storefront: Responsive shopping UI built with React and Tailwind CSS.",
        "education": "B.A. in Interactive Media"
    },
    {
        "name": "Bob React",
        "email": "bob.react@example.com",
        "phone": "+1-555-0144",
        "github": "https://github.com/bobreact",
        "format": "docx",
        "skills": "React, Redux, JavaScript, CSS, Figma",
        "experience": "UI Developer Intern. Converted Figma mockups to React components.",
        "projects": "SaaS Dashboard: React dashboard with Redux state management.",
        "education": "B.S. in Web Development"
    },
    {
        "name": "Charlie Vue",
        "email": "charlie.vue@example.com",
        "phone": "+1-555-0145",
        "github": "https://github.com/charlievue",
        "format": "txt",
        "skills": "Vue.js, Nuxt.js, JavaScript, HTML, CSS, Node.js",
        "experience": "Web Developer Intern.",
        "projects": "Portfolio Platform: Nuxt.js website with Node.js backend.",
        "education": "B.S. in Design & Tech"
    },

    # Additional candidates to reach 50 resumes total
    {
        "name": "Kavita Nair",
        "email": "kavita.nair@example.com",
        "phone": "+1-555-0146",
        "github": "https://github.com/kavitanair",
        "format": "pdf",
        "skills": "Python, LangGraph, FastAPI, ChromaDB, Docker, GCP, PostgreSQL",
        "experience": "AI Intern. Created stateful agent pipelines with LangGraph and ChromaDB.",
        "projects": "Multi-Agent Document QA: Stateful LangGraph workflow with Docker container.",
        "education": "B.Tech in Computer Science"
    },
    {
        "name": "Leo Martinez",
        "email": "leo.martinez@example.com",
        "phone": "+1-555-0147",
        "github": "https://github.com/leomartinez",
        "format": "pdf",
        "skills": "Python, FastAPI, LlamaIndex, Qdrant, Docker, PostgreSQL",
        "experience": "SWE Intern. Built vector search pipelines in Python.",
        "projects": "Enterprise RAG: LlamaIndex search over PDF manuals with Qdrant.",
        "education": "B.S. in Computer Science"
    },
    {
        "name": "Natasha Romanoff",
        "email": "natasha.r@example.com",
        "phone": "+1-555-0148",
        "github": "https://github.com/natasharomanoff",
        "format": "docx",
        "skills": "Python, Django, Celery, Redis, MySQL",
        "experience": "Backend Developer Intern. Built asynchronous background worker tasks in Python.",
        "projects": "Notification Service: High volume email delivery backend with Celery.",
        "education": "B.S. in Computer Science"
    },
    {
        "name": "Ethan Hunt",
        "email": "ethan.hunt@example.com",
        "phone": "+1-555-0149",
        "github": "https://github.com/ethanhunt",
        "format": "pdf",
        "skills": "PHP, Laravel, MySQL, JavaScript",
        "experience": "Web Developer. Maintained Laravel web portals.",
        "projects": "Real Estate CMS: Portal built with PHP Laravel.",
        "education": "B.S. in Information Systems"
    }
]


def create_pdf_resume(cand, file_path: Path):
    doc = SimpleDocTemplate(str(file_path), pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle("NameHeader", parent=styles["Heading1"], fontSize=18, leading=22, spaceAfter=4)
    contact_style = ParagraphStyle("Contact", parent=styles["Normal"], fontSize=9, leading=12, textColor="grey", spaceAfter=14)
    section_style = ParagraphStyle("Section", parent=styles["Heading2"], fontSize=13, leading=16, spaceBefore=8, spaceAfter=4)
    body_style = ParagraphStyle("Body", parent=styles["Normal"], fontSize=10, leading=14, spaceAfter=6)

    story.append(Paragraph(cand["name"], title_style))
    contact_text = f"Email: {cand['email']} | Phone: {cand['phone']} | GitHub: {cand.get('github', 'N/A')}"
    story.append(Paragraph(contact_text, contact_style))

    story.append(Paragraph("Technical Skills", section_style))
    story.append(Paragraph(cand["skills"], body_style))

    story.append(Paragraph("Professional Experience", section_style))
    story.append(Paragraph(cand["experience"], body_style))

    story.append(Paragraph("Key Projects", section_style))
    story.append(Paragraph(cand["projects"], body_style))

    story.append(Paragraph("Education", section_style))
    story.append(Paragraph(cand["education"], body_style))

    doc.build(story)


def create_docx_resume(cand, file_path: Path):
    doc = Document()
    doc.add_heading(cand["name"], level=1)
    p = doc.add_paragraph(f"Email: {cand['email']} | Phone: {cand['phone']} | GitHub: {cand.get('github', 'N/A')}")
    doc.add_heading("Technical Skills", level=2)
    doc.add_paragraph(cand["skills"])
    doc.add_heading("Professional Experience", level=2)
    doc.add_paragraph(cand["experience"])
    doc.add_heading("Projects", level=2)
    doc.add_paragraph(cand["projects"])
    doc.add_heading("Education", level=2)
    doc.add_paragraph(cand["education"])
    doc.save(str(file_path))


def create_txt_resume(cand, file_path: Path):
    content = f"""{cand['name']}
Email: {cand['email']}
Phone: {cand['phone']}
GitHub: {cand.get('github', 'N/A')}

TECHNICAL SKILLS:
{cand['skills']}

WORK EXPERIENCE:
{cand['experience']}

PROJECTS:
{cand['projects']}

EDUCATION:
{cand['education']}
"""
    file_path.write_text(content, encoding="utf-8")


def main():
    print(f"Generating synthetic resumes in {RESUMES_DIR}...")
    for idx, cand in enumerate(CANDIDATES, start=1):
        clean_name = cand["name"].replace(" ", "_").replace("'", "")
        fmt = cand["format"]
        filename = f"candidate_{idx:02d}_{clean_name}.{fmt}"
        target_path = RESUMES_DIR / filename

        if fmt == "pdf":
            create_pdf_resume(cand, target_path)
        elif fmt == "docx":
            create_docx_resume(cand, target_path)
        elif fmt == "txt":
            create_txt_resume(cand, target_path)

    # Add 50th file: A malformed / corrupt file to verify that the batch never crashes!
    corrupt_file = RESUMES_DIR / "candidate_50_corrupted_file.pdf"
    corrupt_file.write_bytes(b"%PDF-1.4\nINVALID CORRUPT BYTES THAT CANNOT BE PARSED\n")

    print(f"Successfully generated {len(list(RESUMES_DIR.iterdir()))} synthetic candidate resumes!")


if __name__ == "__main__":
    main()
